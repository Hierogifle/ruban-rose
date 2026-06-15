# -*- coding: utf-8 -*-
"""Ajoute 2 diapos dédiées aux figures d'Emile (étude binaire) :
 - matrices de confusion (ViT-Small vs ConvNeXt-Tiny)
 - résultats par sous-type & par grossissement
Insérées juste après la diapo 'Matrices de confusion' du collègue (index 12)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

SRC = "Soutenance_Ruban_Rose_2.pptx"

PINK_DARK = RGBColor(0xC2, 0x18, 0x5B)
PINK      = RGBColor(0xF4, 0x8F, 0xB1)
PINK_BG   = RGBColor(0xFC, 0xE4, 0xEC)
INK       = RGBColor(0x1A, 0x1A, 0x2E)
GREY      = RGBColor(0x9E, 0x9E, 0x9E)
GREY_D    = RGBColor(0x42, 0x42, 0x42)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
FONT      = "Calibri"

prs = Presentation(SRC)
LAYOUT = prs.slides[10].slide_layout


def _run(r, text, size, bold, color):
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color


def add_text(slide, l, t, w, h, lines, size=11.5, bold=False, color=INK,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space_after=4):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    if isinstance(lines, str):
        lines = [(lines, {})]
    for i, item in enumerate(lines):
        txt, kw = item if isinstance(item, tuple) else (item, {})
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = kw.get("align", align)
        p.space_after = Pt(kw.get("space_after", space_after))
        p.space_before = Pt(0)
        _run(p.add_run(), txt, kw.get("size", size), kw.get("bold", bold),
             kw.get("color", color))
    return tb


def add_rect(slide, l, t, w, h, fill, rounded=False):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(l), Inches(t), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    if rounded:
        try:
            shp.adjustments[0] = 0.08
        except Exception:
            pass
    return shp


def header(slide, title, subtitle):
    add_rect(slide, 0, 0, 0.18, 5.625, PINK_DARK)
    add_text(slide, 0.35, 0.25, 9.30, 0.70, title, size=28, bold=True, color=INK,
             anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, 0.35, 0.90, 9.30, 0.35, subtitle, size=13, color=GREY)
    add_rect(slide, 0.35, 1.18, 9.30, 0.03, PINK)


def add_image(slide, path, l, t, max_w, max_h, ratio):
    w = max_w
    h = w / ratio
    if h > max_h:
        h = max_h
        w = h * ratio
    left = l + (max_w - w) / 2
    top = t + (max_h - h) / 2
    slide.shapes.add_picture(path, Inches(left), Inches(top), Inches(w), Inches(h))


def note(slide, l, t, w, h, head, body):
    add_rect(slide, l, t, w, h, PINK_BG, rounded=True)
    add_rect(slide, l, t + 0.10, 0.07, h - 0.20, PINK_DARK)
    add_text(slide, l + 0.22, t + 0.06, w - 0.40, h - 0.12, [
        (head, {"size": 11.5, "bold": True, "color": PINK_DARK, "space_after": 3}),
        (body, {"size": 10.5, "color": GREY_D}),
    ], anchor=MSO_ANCHOR.MIDDLE)


# ── Slide 1 : matrices de confusion ──────────────────────────────────────────
s1 = prs.slides.add_slide(LAYOUT)
header(s1, "Matrices de confusion — étude binaire (Emile)",
       "ViT-Small vs ConvNeXt-Tiny · set de test (1 320 images) · Bénin / Malin")
add_image(s1, "results/confusion_matrices.png", 0.70, 1.45, 8.60, 3.05, ratio=2.408)
note(s1, 0.35, 4.62, 9.30, 0.78, "Lecture",
     "Les deux modèles détectent quasi tous les malins (sensibilité élevée). L'écart se joue sur les bénins : ViT-Small en classe correctement bien plus (spécificité 84 % vs 63 %), donc moins de faux positifs que ConvNeXt-Tiny.")

# ── Slide 2 : résultats par sous-type & grossissement ────────────────────────
s2 = prs.slides.add_slide(LAYOUT)
header(s2, "Résultats par sous-type — étude binaire (Emile)",
       "Accuracy par sous-type (sensibilité si malin · spécificité si bénin) et par grossissement · test")
add_image(s2, "results/per_subtype_performance.png", 0.35, 1.40, 5.55, 3.25, ratio=1.678)
add_image(s2, "results/per_magnification_performance.png", 6.05, 1.60, 3.55, 2.55, ratio=1.614)
note(s2, 0.35, 4.72, 9.30, 0.70, "Lecture",
     "tubular_adenoma reste le sous-type difficile pour les deux modèles. ViT-Small domine sur les bénins ; ConvNeXt-Tiny est un peu meilleur sur mucinous/papillary. Performance stable de 40× à 400×.")

# ── Insertion après la diapo 12 (matrices de confusion du collègue) ──────────
sldIdLst = prs.slides._sldIdLst
ids = list(sldIdLst)
new_ids = ids[-2:]
for nid in new_ids:
    sldIdLst.remove(nid)
anchor = list(sldIdLst)[13]      # diapo actuellement en position 13 (Impact WRS)
for nid in new_ids:
    anchor.addprevious(nid)

try:
    prs.save(SRC)
    print("OK — 2 diapos figures ajoutées (positions 13-14). Total:", len(prs.slides._sldIdLst))
except PermissionError:
    TMP = "Soutenance_Ruban_Rose_2.NEW.pptx"
    prs.save(TMP)
    print(f"LOCKED — fichier principal verrouillé (PowerPoint ouvert). Sauvegardé dans {TMP}. Total:", len(prs.slides._sldIdLst))
