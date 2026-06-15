# -*- coding: utf-8 -*-
"""Ajoute une diapo 'Traitement du déséquilibre des classes' juste après
la diapo 6 (Pipeline de traitement) du deck final."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

SRC = "Soutenance_Ruban_Rose_FINAL.pptx"

PINK_DARK = RGBColor(0xC2, 0x18, 0x5B)
PINK      = RGBColor(0xF4, 0x8F, 0xB1)
PINK_BG   = RGBColor(0xFC, 0xE4, 0xEC)
INK       = RGBColor(0x1A, 0x1A, 0x2E)
GREY      = RGBColor(0x9E, 0x9E, 0x9E)
GREY_D    = RGBColor(0x42, 0x42, 0x42)
CARD_BG   = RGBColor(0xFA, 0xFA, 0xFA)
CARD_BD   = RGBColor(0xBD, 0xBD, 0xBD)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
FONT      = "Calibri"

prs = Presentation(SRC)
LAYOUT = prs.slides[10].slide_layout


def _run(r, text, size, bold, color):
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color


def add_text(slide, l, t, w, h, lines, size=11.5, bold=False, color=INK,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space_after=4):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    if isinstance(lines, str):
        lines = [(lines, {})]
    for i, item in enumerate(lines):
        txt, kw = item if isinstance(item, tuple) else (item, {})
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = kw.get("align", align)
        p.space_after = Pt(kw.get("space_after", space_after))
        p.space_before = Pt(0)
        _run(p.add_run(), txt, kw.get("size", size), kw.get("bold", bold),
             kw.get("color", color))
    return tb


def add_rect(slide, l, t, w, h, fill, line=None, rounded=False):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(l), Inches(t), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Emu(9525)
    shp.shadow.inherit = False
    if rounded:
        try:
            shp.adjustments[0] = 0.06
        except Exception:
            pass
    return shp


def card(slide, l, t, w, h, title, accent, lines):
    add_rect(slide, l, t, w, h, CARD_BG, line=CARD_BD, rounded=True)
    add_rect(slide, l, t + 0.10, 0.07, h - 0.20, accent)
    add_text(slide, l + 0.22, t + 0.14, w - 0.40, 0.42, title, size=13.5,
             bold=True, color=PINK_DARK)
    add_text(slide, l + 0.24, t + 0.66, w - 0.46, h - 0.80, lines,
             size=11, color=GREY_D, space_after=9)


# ── nouvelle diapo ───────────────────────────────────────────────────────────
s = prs.slides.add_slide(LAYOUT)
add_rect(s, 0, 0, 0.18, 5.625, PINK_DARK)
add_text(s, 0.35, 0.25, 9.30, 0.70, "Traitement du déséquilibre des classes",
         size=28, bold=True, color=INK, anchor=MSO_ANCHOR.MIDDLE)
add_text(s, 0.35, 0.90, 9.30, 0.35,
         "Deux niveaux de déséquilibre · des stratégies complémentaires",
         size=13, color=GREY)
add_rect(s, 0.35, 1.18, 9.30, 0.03, PINK)

card(s, 0.35, 1.40, 4.55, 3.05, "Niveau 1 — Bénin / Malin (binaire)", PINK_DARK, [
    "Train : 1 695 bénins  vs  3 831 malins  (31 % / 69 %)",
    "→ Poids de classe dans la perte  CrossEntropyLoss(weight=…)",
    "Bénin ×1.63  ·  Malin ×0.72   (sklearn « balanced »)",
    "Une erreur sur un bénin coûte ~2,3× plus → bride le biais « tout malin », améliore la spécificité",
])

card(s, 5.10, 1.40, 4.55, 3.05, "Niveau 2 — Sous-types (8 classes)", PINK, [
    "Forte hétérogénéité : ductal majoritaire  vs  tubular ~4 %, papillary ~8 %",
    "Équipe : WeightedRandomSampler → tire les sous-types rares plus souvent (Tubular vu ~14×)",
    "Emile : augmentation « pyramide » à la volée → léger (dominants) → intensif (rares)",
])

# bandeau clé
add_rect(s, 0.35, 4.60, 9.30, 0.80, PINK_BG, rounded=True)
add_rect(s, 0.35, 4.70, 0.07, 0.60, PINK_DARK)
add_text(s, 0.58, 4.64, 8.95, 0.72, [
    ("💡  Deux leviers distincts", {"size": 11.5, "bold": True, "color": PINK_DARK, "space_after": 2}),
    ("Les poids de classe modifient le COÛT des erreurs (loss) ; le sampler et l'augmentation modifient les DONNÉES vues à l'entraînement. Combinés, ils empêchent le modèle de se rabattre sur la classe majoritaire.",
     {"size": 10.5, "color": GREY_D}),
], anchor=MSO_ANCHOR.MIDDLE)

# ── insertion après la diapo 6 (index 5 = 'Pipeline de traitement') ──────────
sldIdLst = prs.slides._sldIdLst
new_id = list(sldIdLst)[-1]
sldIdLst.remove(new_id)
anchor = list(sldIdLst)[6]      # diapo actuellement en position 6 (Choix métriques P1)
anchor.addprevious(new_id)

try:
    prs.save(SRC)
    print("OK — diapo insérée en position 7. Total:", len(prs.slides._sldIdLst))
except PermissionError:
    TMP = "Soutenance_Ruban_Rose_FINAL.NEW.pptx"
    prs.save(TMP)
    print(f"LOCKED — fichier verrouillé. Sauvegardé dans {TMP}. Total:", len(prs.slides._sldIdLst))
